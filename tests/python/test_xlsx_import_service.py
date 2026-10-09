# tests/python/test_xlsx_import_service.py

from __future__ import annotations

import json
import sqlite3
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from koa_mediatheque.services.xlsx_import_service import (
    apply_xlsx_import,
    preview_xlsx_import,
)


PROTECTED_FIELDS = {
    "media_uuid",
    "version_uuid",
    "filename",
    "original_path",
    "storage_path",
    "sha256",
    "filesize",
    "mimetype",
    "updated_at",
}


LIBRARY_HEADERS = [
    "xlsx_action",
    "media_uuid",
    "version_uuid",
    "title",
    "subtitle",
    "description",
    "summary",
    "original_path",
    "storage_path",
    "filename",
    "extension",
    "mimetype",
    "filesize",
    "sha256",
    "filearea",
    "media_type",
    "language",
    "library_scope",
    "uckk_relevance",
    "target_system",
    "target_export_allowed",
    "public_state",
    "visibility",
    "access_level",
    "ownership_scope",
    "source_type",
    "source_ownership",
    "rights_status",
    "rights_note",
    "restriction_state",
    "restriction_reason",
    "redaction_required",
    "status",
    "provenance",
    "ai_validation_state",
    "ai_confidence",
    "canonical_validation_state",
    "human_review_required",
    "review_queue",
    "review_reason",
    "collections",
    "tags",
    "relations",
    "content_flags",
    "audience_suitability",
    "export_to_uckk",
    "export_to_public",
    "export_policy_note",
    "import_batch",
    "notes",
    "created_at",
    "updated_at",
]


@pytest.fixture()
def connection() -> Iterator[sqlite3.Connection]:
    db = sqlite3.connect(":memory:")
    db.row_factory = sqlite3.Row
    _create_test_schema(db)

    try:
        yield db
    finally:
        db.close()


def _create_test_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE library_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            media_uuid TEXT NOT NULL,
            version_uuid TEXT NOT NULL UNIQUE,
            title TEXT NOT NULL,
            subtitle TEXT,
            description TEXT,
            summary TEXT,
            original_path TEXT NOT NULL,
            storage_path TEXT,
            filename TEXT NOT NULL,
            extension TEXT,
            mimetype TEXT,
            filesize INTEGER,
            sha256 TEXT,
            filearea TEXT DEFAULT 'media_original',
            media_type TEXT DEFAULT 'document',
            language TEXT DEFAULT 'fr',
            library_scope TEXT DEFAULT 'koa',
            uckk_relevance TEXT DEFAULT 'unknown',
            target_system TEXT DEFAULT 'none',
            target_export_allowed INTEGER DEFAULT 0,
            public_state TEXT DEFAULT 'unknown',
            visibility TEXT DEFAULT 'private',
            access_level TEXT DEFAULT 'private',
            ownership_scope TEXT DEFAULT 'unknown',
            source_type TEXT DEFAULT 'unknown',
            source_ownership TEXT DEFAULT 'unknown_source',
            rights_status TEXT DEFAULT 'unknown',
            rights_note TEXT,
            restriction_state TEXT DEFAULT 'none',
            restriction_reason TEXT,
            redaction_required INTEGER DEFAULT 0,
            status TEXT DEFAULT 'active',
            provenance TEXT DEFAULT 'ai_assisted',
            ai_validation_state TEXT DEFAULT 'ai_uncertain',
            ai_confidence REAL,
            canonical_validation_state TEXT DEFAULT 'unverified',
            human_review_required INTEGER DEFAULT 0,
            review_queue TEXT,
            review_reason TEXT,
            collections_json TEXT DEFAULT '[]',
            tags_json TEXT DEFAULT '[]',
            relations_json TEXT DEFAULT '[]',
            content_flags_json TEXT DEFAULT '[]',
            audience_suitability TEXT DEFAULT 'unknown',
            export_to_uckk TEXT DEFAULT 'no',
            export_to_public TEXT DEFAULT 'no',
            export_policy_note TEXT,
            import_batch TEXT,
            notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE xlsx_import_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            import_uuid TEXT,
            xlsx_path TEXT,
            mode TEXT,
            rows_total INTEGER,
            rows_update INTEGER,
            rows_new INTEGER,
            rows_archive INTEGER,
            rows_ignore INTEGER,
            rows_blocked INTEGER,
            status TEXT,
            report_json TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            action TEXT NOT NULL,
            entity_type TEXT,
            entity_uuid TEXT,
            before_json TEXT,
            after_json TEXT,
            actor TEXT,
            note TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        """
    )
    connection.commit()


def insert_library_row(connection: sqlite3.Connection, **overrides: Any) -> dict[str, Any]:
    row = {
        "media_uuid": "11111111-1111-1111-1111-111111111111",
        "version_uuid": "22222222-2222-2222-2222-222222222222",
        "title": "Document original",
        "subtitle": None,
        "description": "Description originale",
        "summary": "Résumé original",
        "original_path": "03_IMPORTS/pending_review/document-original.pdf",
        "storage_path": (
            "02_STORAGE/media_original/"
            "22222222-2222-2222-2222-222222222222_document-original.pdf"
        ),
        "filename": "document-original.pdf",
        "extension": ".pdf",
        "mimetype": "application/pdf",
        "filesize": 12345,
        "sha256": "a" * 64,
        "filearea": "media_original",
        "media_type": "pdf",
        "language": "fr",
        "library_scope": "koa",
        "uckk_relevance": "unknown",
        "target_system": "none",
        "target_export_allowed": 0,
        "public_state": "unknown",
        "visibility": "private",
        "access_level": "private",
        "ownership_scope": "unknown",
        "source_type": "unknown",
        "source_ownership": "unknown_source",
        "rights_status": "unknown",
        "rights_note": "À réviser.",
        "restriction_state": "none",
        "restriction_reason": None,
        "redaction_required": 0,
        "status": "active",
        "provenance": "ai_assisted",
        "ai_validation_state": "ai_uncertain",
        "ai_confidence": 0.5,
        "canonical_validation_state": "unverified",
        "human_review_required": 1,
        "review_queue": "rights_review",
        "review_reason": "Rights/source unknown.",
        "collections_json": json.dumps(["collection_original"], ensure_ascii=False),
        "tags_json": json.dumps(["tag_original"], ensure_ascii=False),
        "relations_json": json.dumps([], ensure_ascii=False),
        "content_flags_json": json.dumps(["copyright_uncertain"], ensure_ascii=False),
        "audience_suitability": "unknown",
        "export_to_uckk": "no",
        "export_to_public": "no",
        "export_policy_note": "Review required before export.",
        "import_batch": "pytest-original",
        "notes": "Original test row.",
        "created_at": "2026-06-15T10:00:00Z",
        "updated_at": "2026-06-15T10:00:00Z",
    }
    row.update(overrides)

    columns = list(row)
    placeholders = ", ".join("?" for _ in columns)

    connection.execute(
        f"""
        INSERT INTO library_rows ({", ".join(columns)})
        VALUES ({placeholders})
        """,
        [row[column] for column in columns],
    )
    connection.commit()

    return row


def get_library_row(connection: sqlite3.Connection, version_uuid: str) -> dict[str, Any]:
    row = connection.execute(
        "SELECT * FROM library_rows WHERE version_uuid = ?",
        (version_uuid,),
    ).fetchone()

    assert row is not None
    return dict(row)


def count_rows(connection: sqlite3.Connection, table: str) -> int:
    row = connection.execute(
        f"SELECT COUNT(*) AS count_value FROM {table}"
    ).fetchone()

    assert row is not None
    return int(row["count_value"])


def make_xlsx_row(**overrides: Any) -> dict[str, Any]:
    row = {
        "xlsx_action": "update",
        "media_uuid": "11111111-1111-1111-1111-111111111111",
        "version_uuid": "22222222-2222-2222-2222-222222222222",
        "title": "Document modifié",
        "subtitle": "Sous-titre modifié",
        "description": "Description modifiée",
        "summary": "Résumé modifié",
        "original_path": "03_IMPORTS/pending_review/changed-path.pdf",
        "storage_path": "02_STORAGE/media_original/changed-storage-path.pdf",
        "filename": "changed-filename.pdf",
        "extension": ".pdf",
        "mimetype": "application/changed",
        "filesize": 99999,
        "sha256": "b" * 64,
        "filearea": "media_original",
        "media_type": "pdf",
        "language": "fr",
        "library_scope": "koa",
        "uckk_relevance": "uckk_reference",
        "target_system": "uckkarchive",
        "target_export_allowed": 0,
        "public_state": "non_public",
        "visibility": "private",
        "access_level": "private",
        "ownership_scope": "unknown",
        "source_type": "unknown",
        "source_ownership": "unknown_source",
        "rights_status": "unknown",
        "rights_note": "Toujours à réviser.",
        "restriction_state": "none",
        "restriction_reason": None,
        "redaction_required": 0,
        "status": "active",
        "provenance": "ai_assisted",
        "ai_validation_state": "ai_classified_needs_review",
        "ai_confidence": 0.77,
        "canonical_validation_state": "unverified",
        "human_review_required": 1,
        "review_queue": "rights_review",
        "review_reason": "Still needs review.",
        "collections": "collection_a; collection_b",
        "tags": "tag_a; tag_b",
        "relations": "",
        "content_flags": "copyright_uncertain",
        "audience_suitability": "unknown",
        "export_to_uckk": "no",
        "export_to_public": "no",
        "export_policy_note": "No export.",
        "import_batch": "pytest-xlsx",
        "notes": "Updated from XLSX.",
        "created_at": "2099-01-01T00:00:00Z",
        "updated_at": "2099-01-01T00:00:00Z",
    }
    row.update(overrides)
    return row


def write_library_xlsx(
    path: Path,
    rows: list[dict[str, Any]],
    *,
    include_library: bool = True,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    if include_library:
        worksheet = workbook.create_sheet("Library")
        worksheet.append(LIBRARY_HEADERS)

        for row in rows:
            worksheet.append([row.get(header) for header in LIBRARY_HEADERS])

    lists_sheet = workbook.create_sheet("Lists")
    lists_sheet.append(["list_name", "value"])
    lists_sheet.append(["MEDIA_STATUS_VALUES", "active"])

    import_report_sheet = workbook.create_sheet("Import_Report")
    import_report_sheet.append(["key", "value"])
    import_report_sheet.append(["source", "pytest"])

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def test_preview_xlsx_import_counts_update_new_archive_ignore_and_blocked_rows(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    xlsx_path = tmp_path / "import.xlsx"
    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(xlsx_action="update"),
            make_xlsx_row(
                xlsx_action="new",
                media_uuid="33333333-3333-3333-3333-333333333333",
                version_uuid="44444444-4444-4444-4444-444444444444",
                title="Nouveau document",
                original_path="03_IMPORTS/pending_review/new.pdf",
                storage_path="02_STORAGE/media_original/new.pdf",
                filename="new.pdf",
                sha256="c" * 64,
            ),
            make_xlsx_row(xlsx_action="archive"),
            make_xlsx_row(xlsx_action="ignore"),
            make_xlsx_row(xlsx_action="delete"),
        ],
    )

    preview = preview_xlsx_import(connection, xlsx_path)

    assert preview.rows_total == 5
    assert preview.rows_update == 1
    assert preview.rows_new == 1
    assert preview.rows_archive == 1
    assert preview.rows_ignore == 1
    assert preview.rows_blocked == 1
    assert preview.errors


def test_preview_xlsx_import_rejects_missing_library_sheet(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    xlsx_path = tmp_path / "missing-library.xlsx"
    write_library_xlsx(xlsx_path, [], include_library=False)

    preview = preview_xlsx_import(connection, xlsx_path)

    assert preview.rows_total == 0
    assert preview.rows_blocked == 1
    assert preview.errors


def test_preview_xlsx_import_rejects_invalid_action(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    xlsx_path = tmp_path / "invalid-action.xlsx"
    write_library_xlsx(xlsx_path, [make_xlsx_row(xlsx_action="delete")])

    preview = preview_xlsx_import(connection, xlsx_path)

    assert preview.rows_total == 1
    assert preview.rows_blocked == 1
    assert preview.errors


def test_preview_xlsx_import_requires_version_uuid_except_new_rows(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    xlsx_path = tmp_path / "missing-version-uuid.xlsx"
    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(xlsx_action="update", version_uuid=""),
            make_xlsx_row(xlsx_action="archive", version_uuid=""),
            make_xlsx_row(xlsx_action="ignore", version_uuid=""),
            make_xlsx_row(
                xlsx_action="new",
                media_uuid="33333333-3333-3333-3333-333333333333",
                version_uuid="44444444-4444-4444-4444-444444444444",
                title="Nouveau document",
                original_path="03_IMPORTS/pending_review/new.pdf",
                storage_path="02_STORAGE/media_original/new.pdf",
                filename="new.pdf",
                sha256="c" * 64,
            ),
        ],
    )

    preview = preview_xlsx_import(connection, xlsx_path)

    assert preview.rows_total == 4
    assert preview.rows_update == 0
    assert preview.rows_archive == 0
    assert preview.rows_new == 1
    assert preview.rows_ignore == 1
    assert preview.rows_blocked == 2


def test_apply_xlsx_import_updates_editable_fields(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    original = insert_library_row(connection)

    xlsx_path = tmp_path / "update.xlsx"
    backup_dir = tmp_path / "backups"
    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(
                xlsx_action="update",
                version_uuid=original["version_uuid"],
                title="Titre modifié",
                description="Description modifiée",
                tags="tag_x; tag_y",
                collections="collection_x; collection_y",
                content_flags="requires_context",
                notes="Notes modifiées.",
            ),
        ],
    )

    result = apply_xlsx_import(connection, xlsx_path, backup_dir=backup_dir)

    assert result.success is True
    assert result.operation == "apply_xlsx_import"
    assert result.result == "imported"

    updated = get_library_row(connection, original["version_uuid"])

    assert updated["title"] == "Titre modifié"
    assert updated["description"] == "Description modifiée"
    assert json.loads(updated["tags_json"]) == ["tag_x", "tag_y"]
    assert json.loads(updated["collections_json"]) == ["collection_x", "collection_y"]
    assert json.loads(updated["content_flags_json"]) == ["requires_context"]
    assert updated["notes"] == "Notes modifiées."


def test_apply_xlsx_import_does_not_update_protected_fields_in_normal_mode(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    original = insert_library_row(connection)

    xlsx_path = tmp_path / "protected-fields.xlsx"
    backup_dir = tmp_path / "backups"
    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(
                xlsx_action="update",
                version_uuid=original["version_uuid"],
                media_uuid="99999999-9999-9999-9999-999999999999",
                filename="changed.pdf",
                original_path="03_IMPORTS/pending_review/changed.pdf",
                storage_path="02_STORAGE/media_original/changed.pdf",
                sha256="f" * 64,
                filesize=999999,
                mimetype="application/changed",
                updated_at="2099-01-01T00:00:00Z",
            ),
        ],
    )

    result = apply_xlsx_import(
        connection,
        xlsx_path,
        backup_dir=backup_dir,
        mode="normal",
    )

    assert result.success is True
    assert result.warnings

    updated = get_library_row(connection, original["version_uuid"])

    for field_name in PROTECTED_FIELDS:
        assert updated[field_name] == original[field_name]


def test_apply_xlsx_import_archives_existing_row(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    original = insert_library_row(connection, status="active")

    xlsx_path = tmp_path / "archive.xlsx"
    backup_dir = tmp_path / "backups"
    write_library_xlsx(
        xlsx_path,
        [make_xlsx_row(xlsx_action="archive", version_uuid=original["version_uuid"])],
    )

    result = apply_xlsx_import(connection, xlsx_path, backup_dir=backup_dir)

    assert result.success is True

    archived = get_library_row(connection, original["version_uuid"])
    assert archived["status"] == "archived"


def test_apply_xlsx_import_ignores_ignore_rows(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    original = insert_library_row(connection, title="Original title")

    xlsx_path = tmp_path / "ignore.xlsx"
    backup_dir = tmp_path / "backups"
    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(
                xlsx_action="ignore",
                version_uuid=original["version_uuid"],
                title="Should not be applied",
            ),
        ],
    )

    result = apply_xlsx_import(connection, xlsx_path, backup_dir=backup_dir)

    assert result.success is True

    unchanged = get_library_row(connection, original["version_uuid"])
    assert unchanged["title"] == "Original title"


def test_apply_xlsx_import_inserts_new_row(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    xlsx_path = tmp_path / "new.xlsx"
    backup_dir = tmp_path / "backups"
    new_version_uuid = "44444444-4444-4444-4444-444444444444"

    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(
                xlsx_action="new",
                media_uuid="33333333-3333-3333-3333-333333333333",
                version_uuid=new_version_uuid,
                title="Nouveau document",
                original_path="03_IMPORTS/pending_review/new.pdf",
                storage_path="02_STORAGE/media_original/new.pdf",
                filename="new.pdf",
                sha256="c" * 64,
            ),
        ],
    )

    result = apply_xlsx_import(connection, xlsx_path, backup_dir=backup_dir)

    assert result.success is True

    inserted = get_library_row(connection, new_version_uuid)
    assert inserted["title"] == "Nouveau document"
    assert inserted["filename"] == "new.pdf"
    assert json.loads(inserted["tags_json"]) == ["tag_a", "tag_b"]


def test_apply_xlsx_import_creates_backup_before_import(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    original = insert_library_row(connection)

    xlsx_path = tmp_path / "backup.xlsx"
    backup_dir = tmp_path / "backups"
    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(
                xlsx_action="update",
                version_uuid=original["version_uuid"],
                title="Updated with backup",
            ),
        ],
    )

    result = apply_xlsx_import(connection, xlsx_path, backup_dir=backup_dir)

    assert result.success is True
    assert backup_dir.exists()
    assert any(path.suffix in {".sqlite", ".db"} for path in backup_dir.iterdir())


def test_apply_xlsx_import_writes_import_and_audit_logs(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    original = insert_library_row(connection)

    xlsx_path = tmp_path / "logs.xlsx"
    backup_dir = tmp_path / "backups"
    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(
                xlsx_action="update",
                version_uuid=original["version_uuid"],
                title="Updated with logs",
            ),
        ],
    )

    result = apply_xlsx_import(connection, xlsx_path, backup_dir=backup_dir)

    assert result.success is True
    assert count_rows(connection, "xlsx_import_log") >= 1
    assert count_rows(connection, "audit_log") >= 1


def test_apply_xlsx_import_dry_run_does_not_modify_database(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    original = insert_library_row(connection, title="Original title")

    xlsx_path = tmp_path / "dry-run.xlsx"
    backup_dir = tmp_path / "backups"
    write_library_xlsx(
        xlsx_path,
        [
            make_xlsx_row(
                xlsx_action="update",
                version_uuid=original["version_uuid"],
                title="Dry run title",
            ),
        ],
    )

    result = apply_xlsx_import(
        connection,
        xlsx_path,
        backup_dir=backup_dir,
        mode="dry_run",
    )

    assert result.success is True
    assert result.result == "dry_run"

    unchanged = get_library_row(connection, original["version_uuid"])
    assert unchanged["title"] == "Original title"


def test_apply_xlsx_import_fails_when_preview_has_blocking_errors(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    xlsx_path = tmp_path / "blocking.xlsx"
    backup_dir = tmp_path / "backups"
    write_library_xlsx(xlsx_path, [make_xlsx_row(xlsx_action="delete")])

    result = apply_xlsx_import(connection, xlsx_path, backup_dir=backup_dir)

    assert result.success is False
    assert result.operation == "apply_xlsx_import"
    assert result.errors

    rows = connection.execute("SELECT title FROM library_rows").fetchall()
    assert len(rows) == 1
    assert rows[0]["title"] == "Document original"