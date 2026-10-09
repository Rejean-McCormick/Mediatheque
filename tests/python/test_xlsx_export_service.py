# tests/python/test_xlsx_export_service.py

from __future__ import annotations

import json
import sqlite3
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from koa_mediatheque.services.xlsx_export_service import export_library_to_xlsx


def create_test_connection() -> sqlite3.Connection:
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row

    connection.executescript(
        """
        CREATE TABLE library_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            media_uuid TEXT NOT NULL,
            version_uuid TEXT NOT NULL UNIQUE,
            title TEXT NOT NULL,
            subtitle TEXT,
            description TEXT,
            summary TEXT,
            original_path TEXT NOT NULL,
            storage_path TEXT,
            filename TEXT NOT NULL,
            extension TEXT,
            mimetype TEXT,
            filesize INTEGER,
            sha256 TEXT,
            filearea TEXT DEFAULT 'media_original',
            media_type TEXT DEFAULT 'document',
            language TEXT DEFAULT 'fr',
            library_scope TEXT DEFAULT 'koa',
            uckk_relevance TEXT DEFAULT 'unknown',
            target_system TEXT DEFAULT 'none',
            target_export_allowed INTEGER DEFAULT 0,
            public_state TEXT DEFAULT 'unknown',
            visibility TEXT DEFAULT 'private',
            access_level TEXT DEFAULT 'private',
            ownership_scope TEXT DEFAULT 'unknown',
            source_type TEXT DEFAULT 'unknown',
            source_ownership TEXT DEFAULT 'unknown_source',
            rights_status TEXT DEFAULT 'unknown',
            rights_note TEXT,
            restriction_state TEXT DEFAULT 'none',
            restriction_reason TEXT,
            redaction_required INTEGER DEFAULT 0,
            status TEXT DEFAULT 'active',
            provenance TEXT DEFAULT 'ai_assisted',
            ai_validation_state TEXT DEFAULT 'ai_uncertain',
            ai_confidence REAL,
            canonical_validation_state TEXT DEFAULT 'unverified',
            human_review_required INTEGER DEFAULT 0,
            review_queue TEXT,
            review_reason TEXT,
            collections_json TEXT DEFAULT '[]',
            tags_json TEXT DEFAULT '[]',
            relations_json TEXT DEFAULT '[]',
            content_flags_json TEXT DEFAULT '[]',
            audience_suitability TEXT DEFAULT 'unknown',
            export_to_uckk TEXT DEFAULT 'no',
            export_to_public TEXT DEFAULT 'no',
            export_policy_note TEXT,
            import_batch TEXT,
            notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            action TEXT NOT NULL,
            entity_type TEXT,
            entity_uuid TEXT,
            before_json TEXT,
            after_json TEXT,
            actor TEXT,
            note TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        """
    )

    return connection


@pytest.fixture()
def connection() -> Iterator[sqlite3.Connection]:
    db = create_test_connection()
    try:
        yield db
    finally:
        db.close()


def insert_library_row(
    connection: sqlite3.Connection,
    **overrides: Any,
) -> dict[str, Any]:
    row = {
        "media_uuid": "11111111-1111-4111-8111-111111111111",
        "version_uuid": "22222222-2222-4222-8222-222222222222",
        "title": "Document test",
        "subtitle": None,
        "description": "Description test",
        "summary": "Résumé test",
        "original_path": "03_IMPORTS/pending_review/document-test.pdf",
        "storage_path": (
            "02_STORAGE/media_original/"
            "22222222-2222-4222-8222-222222222222_document-test.pdf"
        ),
        "filename": "document-test.pdf",
        "extension": ".pdf",
        "mimetype": "application/pdf",
        "filesize": 12345,
        "sha256": "a" * 64,
        "filearea": "media_original",
        "media_type": "document",
        "language": "fr",
        "library_scope": "koa",
        "uckk_relevance": "uckk_reference",
        "target_system": "none",
        "target_export_allowed": 0,
        "public_state": "non_public",
        "visibility": "private",
        "access_level": "private",
        "ownership_scope": "unknown",
        "source_type": "unknown",
        "source_ownership": "unknown_source",
        "rights_status": "unknown",
        "rights_note": "À réviser.",
        "restriction_state": "none",
        "restriction_reason": None,
        "redaction_required": 0,
        "status": "active",
        "provenance": "ai_assisted",
        "ai_validation_state": "ai_classified_needs_review",
        "ai_confidence": 0.72,
        "canonical_validation_state": "unverified",
        "human_review_required": 1,
        "review_queue": "rights_review",
        "review_reason": "Rights/source unknown.",
        "collections_json": json.dumps(["collection_test"], ensure_ascii=False),
        "tags_json": json.dumps(["tag_a", "tag_b"], ensure_ascii=False),
        "relations_json": json.dumps(
            ["references:33333333-3333-4333-8333-333333333333"],
            ensure_ascii=False,
        ),
        "content_flags_json": json.dumps(["copyright_uncertain"], ensure_ascii=False),
        "audience_suitability": "unknown",
        "export_to_uckk": "no",
        "export_to_public": "no",
        "export_policy_note": "Review required before export.",
        "import_batch": "pytest",
        "notes": "Test row.",
        "created_at": "2026-06-15T10:00:00Z",
        "updated_at": "2026-06-15T10:00:00Z",
    }

    row.update(overrides)

    columns = list(row)
    placeholders = ", ".join("?" for _ in columns)

    connection.execute(
        f"""
        INSERT INTO library_rows ({", ".join(columns)})
        VALUES ({placeholders})
        """,
        [row[column] for column in columns],
    )
    connection.commit()

    return row


def workbook_sheet_names(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_sheet_rows(path: Path, sheet_name: str) -> list[list[Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def flatten_sheet_values(rows: list[list[Any]]) -> set[str]:
    return {
        str(cell)
        for row in rows
        for cell in row
        if cell is not None and str(cell).strip()
    }


def exported_first_row(path: Path) -> dict[str, Any]:
    rows = read_sheet_rows(path, "Library")

    assert len(rows) >= 2

    return dict(zip(rows[0], rows[1], strict=False))


def assert_export_success(result: Any, output_path: Path) -> None:
    assert result.success is True
    assert result.operation == "export_library_to_xlsx"
    assert result.result == "exported"
    assert output_path.exists()


@pytest.mark.contract
@pytest.mark.xlsx
def test_export_library_to_xlsx_creates_workbook_with_required_sheets(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    output_path = tmp_path / "koa_library_export.xlsx"

    result = export_library_to_xlsx(connection, output_path)

    assert_export_success(result, output_path)

    sheet_names = workbook_sheet_names(output_path)

    assert "Library" in sheet_names
    assert "Lists" in sheet_names
    assert "Import_Report" in sheet_names


@pytest.mark.contract
@pytest.mark.xlsx
def test_export_library_to_xlsx_writes_library_headers_and_row_values(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    inserted = insert_library_row(connection)
    output_path = tmp_path / "koa_library_export.xlsx"

    result = export_library_to_xlsx(connection, output_path)

    assert_export_success(result, output_path)

    rows = read_sheet_rows(output_path, "Library")
    headers = rows[0]
    exported = dict(zip(headers, rows[1], strict=False))

    assert "version_uuid" in headers
    assert "media_uuid" in headers
    assert "title" in headers
    assert "sha256" in headers
    assert "collections" in headers
    assert "tags" in headers
    assert "relations" in headers
    assert "content_flags" in headers

    assert exported["version_uuid"] == inserted["version_uuid"]
    assert exported["media_uuid"] == inserted["media_uuid"]
    assert exported["title"] == inserted["title"]
    assert exported["filename"] == inserted["filename"]
    assert exported["sha256"] == inserted["sha256"]
    assert exported["canonical_validation_state"] == "unverified"
    assert exported["human_review_required"] == 1


@pytest.mark.contract
@pytest.mark.xlsx
def test_export_library_to_xlsx_converts_json_text_fields_to_semicolon_strings(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(
        connection,
        collections_json=json.dumps(
            ["collection_a", "collection_b"],
            ensure_ascii=False,
        ),
        tags_json=json.dumps(
            ["tag_a", "tag_b"],
            ensure_ascii=False,
        ),
        relations_json=json.dumps(
            [
                "belongs_to_collection:aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa",
                "references:bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb",
            ],
            ensure_ascii=False,
        ),
        content_flags_json=json.dumps(
            ["copyright_uncertain", "requires_context"],
            ensure_ascii=False,
        ),
    )

    output_path = tmp_path / "koa_library_export.xlsx"

    result = export_library_to_xlsx(connection, output_path)

    assert_export_success(result, output_path)

    exported = exported_first_row(output_path)

    assert exported["collections"] == "collection_a; collection_b"
    assert exported["tags"] == "tag_a; tag_b"
    assert (
        exported["relations"]
        == "belongs_to_collection:aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa; "
        "references:bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb"
    )
    assert exported["content_flags"] == "copyright_uncertain; requires_context"


@pytest.mark.contract
@pytest.mark.xlsx
def test_export_library_to_xlsx_writes_lists_sheet(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    output_path = tmp_path / "koa_library_export.xlsx"

    result = export_library_to_xlsx(
        connection,
        output_path,
        include_lists=True,
    )

    assert_export_success(result, output_path)

    rows = read_sheet_rows(output_path, "Lists")
    flattened_values = flatten_sheet_values(rows)

    assert "MEDIA_STATUS_VALUES" in flattened_values
    assert "CANONICAL_VALIDATION_VALUES" in flattened_values
    assert "VISIBILITY_VALUES" in flattened_values
    assert "PUBLIC_STATE_VALUES" in flattened_values
    assert "MEDIA_TYPE_VALUES" in flattened_values


@pytest.mark.contract
@pytest.mark.xlsx
def test_export_library_to_xlsx_writes_import_report_sheet(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    output_path = tmp_path / "koa_library_export.xlsx"

    result = export_library_to_xlsx(
        connection,
        output_path,
        include_import_report=True,
    )

    assert_export_success(result, output_path)

    rows = read_sheet_rows(output_path, "Import_Report")
    flattened_values = flatten_sheet_values(rows)

    assert "export_library_to_xlsx" in flattened_values
    assert "row_count" in flattened_values
    assert "1" in flattened_values


@pytest.mark.contract
@pytest.mark.xlsx
def test_export_library_to_xlsx_applies_filters(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(
        connection,
        media_uuid="11111111-1111-4111-8111-000000000001",
        version_uuid="22222222-2222-4222-8222-000000000001",
        title="Active record",
        status="active",
    )
    insert_library_row(
        connection,
        media_uuid="33333333-3333-4333-8333-000000000002",
        version_uuid="44444444-4444-4444-8444-000000000002",
        title="Archived record",
        status="archived",
    )

    output_path = tmp_path / "koa_library_export.xlsx"

    result = export_library_to_xlsx(
        connection,
        output_path,
        filters={"status": "active"},
    )

    assert_export_success(result, output_path)

    rows = read_sheet_rows(output_path, "Library")
    data_rows = rows[1:]

    assert len(data_rows) == 1

    exported = dict(zip(rows[0], data_rows[0], strict=False))
    assert exported["title"] == "Active record"
    assert exported["status"] == "active"


@pytest.mark.contract
@pytest.mark.xlsx
def test_export_library_to_xlsx_creates_parent_directory(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    output_path = tmp_path / "nested" / "exports" / "koa_library_export.xlsx"

    result = export_library_to_xlsx(connection, output_path)

    assert_export_success(result, output_path)
    assert output_path.parent.exists()


@pytest.mark.contract
@pytest.mark.xlsx
def test_export_library_to_xlsx_rejects_non_xlsx_output_path(
    connection: sqlite3.Connection,
    tmp_path: Path,
) -> None:
    insert_library_row(connection)

    output_path = tmp_path / "koa_library_export.csv"

    result = export_library_to_xlsx(connection, output_path)

    assert result.success is False
    assert result.operation == "export_library_to_xlsx"
    assert result.errors