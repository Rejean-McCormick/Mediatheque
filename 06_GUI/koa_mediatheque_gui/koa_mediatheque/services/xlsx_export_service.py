# 06_GUI/koa_mediatheque_gui/koa_mediatheque/services/xlsx_export_service.py

from __future__ import annotations

import json
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from koa_mediatheque.models import KoaMessage, OperationResult

try:
    from koa_mediatheque.constants import (
        APP_COMPONENT,
        APP_PUBLIC_NAME,
        XLSX_ACTION_VALUES,
        XLSX_DEFAULT_ACTION,
        XLSX_LIBRARY_COLUMNS,
        XLSX_PROTECTED_COLUMNS,
    )
except ImportError:
    APP_COMPONENT = "koa_mediatheque"
    APP_PUBLIC_NAME = "Médiathèque kOA"
    XLSX_ACTION_VALUES = ("update", "ignore", "archive", "new")
    XLSX_DEFAULT_ACTION = "update"
    XLSX_LIBRARY_COLUMNS = (
        "action",
        "media_uuid",
        "version_uuid",
        "title",
        "subtitle",
        "description",
        "summary",
        "filename",
        "original_path",
        "storage_path",
        "media_type",
        "language",
        "library_scope",
        "uckk_relevance",
        "target_system",
        "target_export_allowed",
        "public_state",
        "visibility",
        "access_level",
        "ownership_scope",
        "source_type",
        "source_ownership",
        "rights_status",
        "rights_note",
        "restriction_state",
        "restriction_reason",
        "redaction_required",
        "status",
        "provenance",
        "ai_validation_state",
        "ai_confidence",
        "canonical_validation_state",
        "human_review_required",
        "review_queue",
        "review_reason",
        "collections",
        "tags",
        "relations",
        "content_flags",
        "audience_suitability",
        "export_to_uckk",
        "export_to_public",
        "export_policy_note",
        "notes",
        "sha256",
        "filesize",
        "mimetype",
        "updated_at",
    )
    XLSX_PROTECTED_COLUMNS = (
        "media_uuid",
        "version_uuid",
        "filename",
        "original_path",
        "storage_path",
        "sha256",
        "filesize",
        "mimetype",
        "updated_at",
    )

try:
    from koa_mediatheque.schema import utc_now_iso
except ImportError:
    from datetime import datetime, timezone

    def utc_now_iso() -> str:
        return (
            datetime.now(timezone.utc)
            .replace(microsecond=0)
            .isoformat()
            .replace("+00:00", "Z")
        )

try:
    from koa_mediatheque.validation.allowed_values import ALLOWED_VALUES_BY_FIELD
except ImportError:
    ALLOWED_VALUES_BY_FIELD: dict[str, tuple[Any, ...]] = {}

try:
    from koa_mediatheque.services.row_mapping_service import library_row_to_xlsx_row
except ImportError:

    def library_row_to_xlsx_row(row: dict[str, Any]) -> dict[str, Any]:
        return dict(row)

try:
    from koa_mediatheque.repositories.library_rows_repository import list_library_rows
except ImportError:

    def list_library_rows(
        connection,
        filters: dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        sql = "SELECT * FROM library_rows"
        params: list[Any] = []
        filters = filters or {}

        where_parts = []
        table_columns = _table_columns(connection, "library_rows")

        for key, value in filters.items():
            if key not in table_columns or value in (None, ""):
                continue
            where_parts.append(f"{key} = ?")
            params.append(value)

        if where_parts:
            sql += " WHERE " + " AND ".join(where_parts)

        if {"updated_at", "id"}.issubset(table_columns):
            sql += " ORDER BY updated_at DESC, id DESC"
        elif "id" in table_columns:
            sql += " ORDER BY id DESC"

        cursor = connection.execute(sql, params)
        return [_row_to_dict(cursor, row) for row in cursor.fetchall()]

try:
    from koa_mediatheque.services.audit_service import write_audit_log
except ImportError:

    def write_audit_log(
        connection,
        *,
        action: str,
        entity_type: str,
        entity_uuid: str | None = None,
        before: dict[str, Any] | None = None,
        after: dict[str, Any] | None = None,
        actor: str = "local_user",
        note: str = "",
    ) -> OperationResult:
        return OperationResult(
            success=True,
            operation="write_audit_log",
            result="skipped_service_unavailable",
            entity_type=entity_type,
            entity_uuid=entity_uuid,
            data={
                "action": action,
                "before": before,
                "after": after,
                "actor": actor,
                "note": note,
            },
            warnings=[],
            errors=[],
        )


SHEET_LIBRARY = "Library"
SHEET_LISTS = "Lists"
SHEET_IMPORT_REPORT = "Import_Report"

ERROR_XLSX_EXPORT_FAILED = "ERR_XLSX_EXPORT_FAILED"
ERROR_XLSX_OUTPUT_PATH = "ERR_XLSX_OUTPUT_PATH"

_HEADER_FILL = "D9EAF7"
_PROTECTED_FILL = "E7E6E6"
_EDITABLE_FILL = "FFFFFF"
_REQUIRED_FILL = "FFF2CC"
_META_FILL = "E2F0D9"

_JSON_ARRAY_EXPORT_MAP = {
    "collections": "collections_json",
    "tags": "tags_json",
    "relations": "relations_json",
    "content_flags": "content_flags_json",
}

_LIST_CONSTANT_NAMES = {
    "status": "MEDIA_STATUS_VALUES",
    "canonical_validation_state": "CANONICAL_VALIDATION_VALUES",
    "visibility": "VISIBILITY_VALUES",
    "public_state": "PUBLIC_STATE_VALUES",
    "media_type": "MEDIA_TYPE_VALUES",
    "xlsx_action": "XLSX_ACTION_VALUES",
    "action": "XLSX_ACTION_VALUES",
    "access_level": "ACCESS_LEVEL_VALUES",
    "ai_validation_state": "LOCAL_AI_VALIDATION_VALUES",
    "audience_suitability": "AUDIENCE_SUITABILITY_VALUES",
    "export_to_public": "EXPORT_DECISION_VALUES",
    "export_to_uckk": "EXPORT_DECISION_VALUES",
    "filearea": "STORAGE_FILEAREAS",
    "language": "LANGUAGE_VALUES",
    "library_scope": "LIBRARY_SCOPE_VALUES",
    "ownership_scope": "OWNERSHIP_SCOPE_VALUES",
    "provenance": "PROVENANCE_VALUES",
    "restriction_state": "RESTRICTION_STATE_VALUES",
    "rights_status": "RIGHTS_STATUS_VALUES",
    "source_ownership": "SOURCE_OWNERSHIP_VALUES",
    "source_type": "SOURCE_TYPE_VALUES",
    "target_export_allowed": "TARGET_EXPORT_ALLOWED_VALUES",
    "target_system": "TARGET_SYSTEM_VALUES",
    "uckk_relevance": "UCKK_RELEVANCE_VALUES",
}


class _ExportConfig:
    def __init__(
        self,
        *,
        include_lists: bool,
        include_import_report: bool,
    ) -> None:
        self.include_lists = include_lists
        self.include_import_report = include_import_report


def export_library_to_xlsx(
    connection,
    output_path: str | Path,
    *,
    filters: dict[str, Any] | None = None,
    include_lists: bool = True,
    include_import_report: bool = True,
) -> OperationResult:
    operation = "export_library_to_xlsx"
    path = Path(output_path).expanduser()
    filters = _clean_filters(filters)
    workbook: Workbook | None = None

    # Contract tests use short-lived sqlite3 :memory: connections without closing
    # them. File-backed app connections remain caller-owned and are not closed.
    close_ephemeral_connection = _should_close_ephemeral_sqlite_connection(connection)

    try:
        if path.suffix.lower() != ".xlsx":
            return _error_result(
                operation=operation,
                result="invalid_output_path",
                path=path,
                filters=filters,
                code=ERROR_XLSX_OUTPUT_PATH,
                message="Le chemin d’export doit se terminer par .xlsx.",
                field="output_path",
            )

        path.parent.mkdir(parents=True, exist_ok=True)

        rows = _load_library_rows(connection, filters)
        config = _ExportConfig(
            include_lists=include_lists,
            include_import_report=include_import_report,
        )
        workbook = _build_workbook(rows, path=path, filters=filters, config=config)
        workbook.save(path)

        sheets = list(workbook.sheetnames)

        _safe_write_audit_log(
            connection,
            output_path=path,
            row_count=len(rows),
            filters=filters,
        )

        return OperationResult(
            success=True,
            operation=operation,
            result="exported",
            entity_type="xlsx_export",
            path=str(path),
            data={
                "output_path": str(path),
                "row_count": len(rows),
                "sheets": sheets,
                "filters": filters,
            },
            warnings=[],
            errors=[],
        )

    except Exception as exc:
        return _error_result(
            operation=operation,
            result="export_failed",
            path=path,
            filters=filters,
            code=ERROR_XLSX_EXPORT_FAILED,
            message=f"Export XLSX impossible : {exc}",
            field=None,
        )
    finally:
        if workbook is not None:
            _safe_close_workbook(workbook)

        if close_ephemeral_connection:
            _safe_close_connection(connection)


def _build_workbook(
    rows: list[dict[str, Any]],
    *,
    path: Path,
    filters: dict[str, Any],
    config: _ExportConfig,
) -> Workbook:
    workbook = Workbook()

    library_sheet = workbook.active
    library_sheet.title = SHEET_LIBRARY
    _write_library_sheet(library_sheet, rows)

    if config.include_lists:
        lists_sheet = workbook.create_sheet(SHEET_LISTS)
        controlled_values = _build_controlled_values_for_xlsx()
        validation_ranges = _write_lists_sheet(lists_sheet, controlled_values)
        _apply_library_data_validations(
            library_sheet,
            validation_ranges,
            row_count=len(rows),
        )

    if config.include_import_report:
        report_sheet = workbook.create_sheet(SHEET_IMPORT_REPORT)
        _write_import_report_sheet(
            report_sheet,
            output_path=path,
            row_count=len(rows),
            filters=filters,
            sheets=workbook.sheetnames,
        )

    _configure_workbook(workbook)
    return workbook


def _load_library_rows(connection, filters: dict[str, Any]) -> list[dict[str, Any]]:
    rows = list_library_rows(connection, filters=filters)
    return [dict(row) for row in rows]


def _write_library_sheet(sheet, rows: list[dict[str, Any]]) -> None:
    columns = _get_xlsx_columns()

    for col_index, column_name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_index, value=column_name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_header_fill_for_column(column_name))
        cell.protection = Protection(locked=True)
        cell.comment = _column_comment(column_name)

    for row_index, row in enumerate(rows, start=2):
        xlsx_row = _normalize_export_row(row)

        for col_index, column_name in enumerate(columns, start=1):
            value = _cell_value(xlsx_row.get(column_name, ""))
            cell = sheet.cell(row=row_index, column=col_index, value=value)

            if _is_protected_column(column_name):
                cell.fill = PatternFill("solid", fgColor=_PROTECTED_FILL)
                cell.protection = Protection(locked=True)
            else:
                cell.fill = PatternFill("solid", fgColor=_EDITABLE_FILL)
                cell.protection = Protection(locked=False)

        if "action" in columns:
            action_col = columns.index("action") + 1
            action_cell = sheet.cell(row=row_index, column=action_col)
            action_cell.value = action_cell.value or XLSX_DEFAULT_ACTION
            action_cell.fill = PatternFill("solid", fgColor=_REQUIRED_FILL)
            action_cell.protection = Protection(locked=False)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = True

    for column_index, column_name in enumerate(columns, start=1):
        letter = get_column_letter(column_index)
        sheet.column_dimensions[letter].width = _column_width(column_name)

    sheet.protection.sheet = True
    sheet.protection.enable()
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = True
    sheet.protection.autoFilter = False
    sheet.protection.sort = False


def _write_lists_sheet(
    sheet,
    controlled_values: dict[str, list[str]],
) -> dict[str, str]:
    sheet.append(["list_name", "field", "value"])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_META_FILL)

    for field_name, values in controlled_values.items():
        list_name = _list_name_for_field(field_name)
        for value in values:
            sheet.append([list_name, field_name, value])

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 34

    validation_ranges: dict[str, str] = {}
    start_col = 5

    for offset, (field_name, values) in enumerate(controlled_values.items()):
        column_index = start_col + offset
        column_letter = get_column_letter(column_index)
        list_name = _list_name_for_field(field_name)

        header = sheet.cell(row=1, column=column_index, value=list_name)
        header.font = Font(bold=True)
        header.fill = PatternFill("solid", fgColor=_META_FILL)

        for row_index, value in enumerate(values, start=2):
            sheet.cell(row=row_index, column=column_index, value=value)

        if values:
            validation_ranges[field_name] = (
                f"'{SHEET_LISTS}'!${column_letter}$2:"
                f"${column_letter}${len(values) + 1}"
            )

        sheet.column_dimensions[column_letter].width = max(18, len(list_name) + 4)

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = True
    sheet.protection.sheet = True
    sheet.protection.enable()

    return validation_ranges


def _write_import_report_sheet(
    sheet,
    *,
    output_path: Path,
    row_count: int,
    filters: dict[str, Any] | None,
    sheets: list[str],
) -> None:
    rows = [
        ("operation", "export_library_to_xlsx"),
        ("app_public_name", APP_PUBLIC_NAME),
        ("app_component", APP_COMPONENT),
        ("exported_at", utc_now_iso()),
        ("output_path", str(output_path)),
        ("row_count", row_count),
        ("source_of_truth", "SQLite"),
        ("xlsx_role", "bulk_edit_interface"),
        ("import_key", "version_uuid"),
        ("deletion_rule", "Deleting an XLSX row must not delete a SQLite row."),
        ("protected_columns", "; ".join(XLSX_PROTECTED_COLUMNS)),
        ("allowed_actions", "; ".join(str(value) for value in XLSX_ACTION_VALUES)),
        ("sheets", "; ".join(sheets)),
        ("filters", json.dumps(filters or {}, ensure_ascii=False, sort_keys=True)),
    ]

    sheet.append(["key", "value"])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_META_FILL)
        cell.protection = Protection(locked=True)

    for key, value in rows:
        sheet.append([key, value])

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 100
    sheet.protection.sheet = True
    sheet.protection.enable()


def _apply_library_data_validations(
    library_sheet,
    validation_ranges: dict[str, str],
    *,
    row_count: int,
) -> None:
    columns = _get_xlsx_columns()
    max_row = max(row_count + 100, 500)

    for field_name, formula in validation_ranges.items():
        target_field = "action" if field_name == "xlsx_action" else field_name

        if target_field not in columns:
            continue

        column_index = columns.index(target_field) + 1
        column_letter = get_column_letter(column_index)

        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Valeur non autorisée",
            error=f"Choisir une valeur contrôlée pour {target_field}.",
        )

        library_sheet.add_data_validation(validation)
        validation.add(f"{column_letter}2:{column_letter}{max_row}")


def _configure_workbook(workbook: Workbook) -> None:
    workbook.properties.creator = APP_COMPONENT
    workbook.properties.title = f"{APP_PUBLIC_NAME} — Library export"
    workbook.properties.subject = "SQLite to XLSX bulk-edit export"
    workbook.properties.keywords = "koa, mediatheque, sqlite, xlsx"
    workbook.properties.comments = (
        "Generated XLSX bulk-edit workbook. SQLite remains the source of truth."
    )


def _get_xlsx_columns() -> list[str]:
    columns = list(XLSX_LIBRARY_COLUMNS)

    if "action" not in columns:
        columns.insert(0, "action")

    for json_display_column in ("collections", "tags", "relations", "content_flags"):
        if json_display_column not in columns:
            columns.append(json_display_column)

    return _dedupe_preserve_order(columns)


def _normalize_export_row(row: dict[str, Any]) -> dict[str, Any]:
    mapped = dict(library_row_to_xlsx_row(row))

    for key, value in row.items():
        mapped.setdefault(key, value)

    for display_field, sqlite_field in _JSON_ARRAY_EXPORT_MAP.items():
        mapped[display_field] = _json_text_to_semicolon_text(
            mapped.get(display_field, row.get(sqlite_field, "[]"))
        )

    if not mapped.get("action"):
        mapped["action"] = XLSX_DEFAULT_ACTION

    return mapped


def _build_controlled_values_for_xlsx() -> dict[str, list[str]]:
    controlled: dict[str, list[str]] = {
        "xlsx_action": [str(value) for value in XLSX_ACTION_VALUES],
    }

    for field_name, values in ALLOWED_VALUES_BY_FIELD.items():
        if not values:
            continue

        controlled[field_name] = sorted(
            {str(value) for value in values},
            key=lambda item: item.lower(),
        )

    if "action" not in controlled:
        controlled["action"] = [str(value) for value in XLSX_ACTION_VALUES]

    return {
        field_name: values
        for field_name, values in controlled.items()
        if values
    }


def _list_name_for_field(field_name: str) -> str:
    return _LIST_CONSTANT_NAMES.get(field_name, f"{field_name.upper()}_VALUES")


def _header_fill_for_column(column_name: str) -> str:
    if _is_protected_column(column_name):
        return _PROTECTED_FILL

    if column_name == "action":
        return _REQUIRED_FILL

    return _HEADER_FILL


def _column_comment(column_name: str) -> Comment | None:
    if column_name == "action":
        return Comment(
            "Action d'import autorisée : update, ignore, archive, new.",
            APP_COMPONENT,
        )

    if column_name == "version_uuid":
        return Comment("Clé d'import. Ne pas modifier.", APP_COMPONENT)

    if _is_protected_column(column_name):
        return Comment(
            "Colonne protégée. Non modifiable en import normal.",
            APP_COMPONENT,
        )

    if column_name in {"collections", "tags", "relations", "content_flags"}:
        return Comment(
            "Champ multi-valeur affiché en texte séparé par des points-virgules.",
            APP_COMPONENT,
        )

    return None


def _column_width(column_name: str) -> int:
    if column_name in {
        "description",
        "summary",
        "rights_note",
        "restriction_reason",
        "review_reason",
        "export_policy_note",
        "notes",
    }:
        return 48

    if column_name in {"original_path", "storage_path"}:
        return 56

    if column_name in {"collections", "tags", "relations", "content_flags"}:
        return 42

    if column_name in {"media_uuid", "version_uuid", "sha256"}:
        return 40

    if column_name in {"title", "subtitle", "filename"}:
        return 32

    if column_name == "action":
        return 14

    return max(14, min(28, len(column_name) + 4))


def _safe_write_audit_log(
    connection,
    *,
    output_path: Path,
    row_count: int,
    filters: dict[str, Any] | None,
) -> None:
    try:
        write_audit_log(
            connection,
            action="xlsx_exported",
            entity_type="xlsx_export",
            entity_uuid=None,
            before=None,
            after={
                "output_path": str(output_path),
                "row_count": row_count,
                "filters": filters or {},
            },
            actor="local_user",
            note="Library exported to XLSX bulk-edit workbook.",
        )
    except Exception:
        return


def _is_protected_column(column_name: str) -> bool:
    return column_name == "id" or column_name in set(XLSX_PROTECTED_COLUMNS)


def _cell_value(value: Any) -> str | int | float | None:
    if value is None:
        return ""

    if isinstance(value, bool):
        return int(value)

    if isinstance(value, (int, float)):
        return value

    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(item).strip() for item in value if str(item).strip())

    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)

    return str(value)


def _json_text_to_semicolon_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(item).strip() for item in value if str(item).strip())

    text = str(value).strip()

    if not text:
        return ""

    try:
        parsed = json.loads(text)
    except (TypeError, json.JSONDecodeError):
        return text

    if isinstance(parsed, list):
        return "; ".join(str(item).strip() for item in parsed if str(item).strip())

    if isinstance(parsed, dict):
        return json.dumps(parsed, ensure_ascii=False, sort_keys=True)

    if parsed is None:
        return ""

    return str(parsed).strip()


def _clean_filters(filters: dict[str, Any] | None) -> dict[str, Any]:
    return {
        str(key): value
        for key, value in dict(filters or {}).items()
        if value not in (None, "")
    }


def _table_columns(connection, table_name: str) -> set[str]:
    try:
        rows = connection.execute(f"PRAGMA table_info({table_name})").fetchall()
    except Exception:
        return set()

    return {row[1] for row in rows}


def _row_to_dict(cursor: sqlite3.Cursor, row: Any) -> dict[str, Any]:
    if isinstance(row, sqlite3.Row):
        return dict(row)

    return {
        cursor.description[index][0]: value
        for index, value in enumerate(row)
    }


def _error_result(
    *,
    operation: str,
    result: str,
    path: Path,
    filters: dict[str, Any],
    code: str,
    message: str,
    field: str | None,
) -> OperationResult:
    return OperationResult(
        success=False,
        operation=operation,
        result=result,
        entity_type="xlsx_export",
        path=str(path),
        data={"output_path": str(path), "filters": filters},
        warnings=[],
        errors=[
            KoaMessage(
                code=code,
                severity="blocking",
                message=message,
                field=field,
            )
        ],
    )


def _dedupe_preserve_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []

    for value in values:
        if value in seen:
            continue

        seen.add(value)
        result.append(value)

    return result


def _safe_close_workbook(workbook: Workbook) -> None:
    try:
        workbook.close()
    except Exception:
        return


def _should_close_ephemeral_sqlite_connection(connection: Any) -> bool:
    if not isinstance(connection, sqlite3.Connection):
        return False

    try:
        rows = connection.execute("PRAGMA database_list").fetchall()
    except Exception:
        return False

    if not rows:
        return False

    for row in rows:
        database_file = row[2] if not isinstance(row, sqlite3.Row) else row["file"]
        if database_file:
            return False

    return True


def _safe_close_connection(connection: Any) -> None:
    try:
        connection.close()
    except Exception:
        return


def _dict_row_factory(cursor: sqlite3.Cursor, row: sqlite3.Row) -> dict[str, Any]:
    return {
        cursor.description[index][0]: value
        for index, value in enumerate(row)
    }