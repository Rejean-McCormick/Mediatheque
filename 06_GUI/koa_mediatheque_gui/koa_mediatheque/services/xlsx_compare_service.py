# 06_GUI/koa_mediatheque_gui/koa_mediatheque/services/xlsx_compare_service.py

from __future__ import annotations

import json
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from koa_mediatheque.errors import (
    ERR_FILE_NOT_FOUND,
    ERR_PROTECTED_FIELD_UPDATE,
    ERR_VERSION_UUID_MISSING,
    ERR_XLSX_INVALID_ACTION,
    ERR_XLSX_MISSING_SHEET,
)
from koa_mediatheque.models import ImportPreview, KoaMessage
from koa_mediatheque.services.row_mapping_service import xlsx_row_to_update_dict

try:
    from koa_mediatheque.constants import (
        XLSX_PROTECTED_COLUMNS,
        TECHNICAL_PROTECTED_FIELDS,
    )
except ImportError:
    XLSX_PROTECTED_COLUMNS: list[str] = [
        "media_uuid",
        "version_uuid",
        "filename",
        "original_path",
        "storage_path",
        "sha256",
        "filesize",
        "mimetype",
        "updated_at",
    ]
    TECHNICAL_PROTECTED_FIELDS: list[str] = [
        "sha256",
        "filesize",
        "mimetype",
        "filename",
        "extension",
        "original_path",
        "storage_path",
    ]

try:
    from koa_mediatheque.repositories.library_rows_repository import (
        get_library_row_by_version_uuid,
    )
except ImportError:

    def get_library_row_by_version_uuid(connection, version_uuid: str) -> dict[str, Any] | None:
        connection.row_factory = _dict_row_factory
        cursor = connection.execute(
            "SELECT * FROM library_rows WHERE version_uuid = ? LIMIT 1",
            (version_uuid,),
        )
        row = cursor.fetchone()
        return dict(row) if row else None


SHEET_LIBRARY = "Library"
ALLOWED_IMPORT_ACTIONS = {"update", "ignore", "archive", "new"}


def compare_xlsx_to_sqlite(
    connection,
    xlsx_path: str | Path,
) -> ImportPreview:
    """
    Compare one XLSX Library sheet against SQLite library_rows.

    This function is read-only. It does not update SQLite, does not create a
    backup, and does not write import logs. It reports what would differ if the
    workbook were used as an import source.
    """
    import_uuid = str(uuid4())
    source_path = Path(xlsx_path).expanduser()
    warnings: list[KoaMessage] = []
    errors: list[KoaMessage] = []

    if not source_path.exists() or not source_path.is_file():
        return ImportPreview(
            import_uuid=import_uuid,
            source_path=str(source_path),
            rows_total=0,
            rows_update=0,
            rows_new=0,
            rows_archive=0,
            rows_ignore=0,
            rows_blocked=0,
            changes=[],
            warnings=warnings,
            errors=[
                KoaMessage(
                    code=ERR_FILE_NOT_FOUND,
                    severity="blocking",
                    field="xlsx_path",
                    message=f"Fichier XLSX introuvable : {source_path}",
                )
            ],
        )

    try:
        workbook = load_workbook(source_path, data_only=True)
    except Exception as exc:
        return ImportPreview(
            import_uuid=import_uuid,
            source_path=str(source_path),
            rows_total=0,
            rows_update=0,
            rows_new=0,
            rows_archive=0,
            rows_ignore=0,
            rows_blocked=0,
            changes=[],
            warnings=warnings,
            errors=[
                KoaMessage(
                    code="ERR_XLSX_READ_FAILED",
                    severity="blocking",
                    field="xlsx_path",
                    message=f"Lecture XLSX impossible : {exc}",
                )
            ],
        )

    if SHEET_LIBRARY not in workbook.sheetnames:
        return ImportPreview(
            import_uuid=import_uuid,
            source_path=str(source_path),
            rows_total=0,
            rows_update=0,
            rows_new=0,
            rows_archive=0,
            rows_ignore=0,
            rows_blocked=0,
            changes=[],
            warnings=warnings,
            errors=[
                KoaMessage(
                    code=ERR_XLSX_MISSING_SHEET,
                    severity="blocking",
                    field=SHEET_LIBRARY,
                    message="Le classeur XLSX doit contenir une feuille Library.",
                )
            ],
        )

    sheet = workbook[SHEET_LIBRARY]
    rows = _read_library_sheet_rows(sheet)

    changes = [
        _compare_one_row(
            connection,
            row=row,
            row_number=row_number,
        )
        for row_number, row in rows
    ]

    return ImportPreview(
        import_uuid=import_uuid,
        source_path=str(source_path),
        rows_total=len(changes),
        rows_update=_count_changes(changes, "update"),
        rows_new=_count_changes(changes, "new"),
        rows_archive=_count_changes(changes, "archive"),
        rows_ignore=_count_changes(changes, "ignore"),
        rows_blocked=sum(1 for change in changes if change.get("blocked")),
        changes=changes,
        warnings=warnings,
        errors=errors,
    )


def _compare_one_row(
    connection,
    *,
    row: dict[str, Any],
    row_number: int,
) -> dict[str, Any]:
    action = _normalize_action(
        row.get("action") or row.get("import_action") or row.get("_action")
    )
    version_uuid = _clean_string(row.get("version_uuid"))

    if action not in ALLOWED_IMPORT_ACTIONS:
        return _blocked_change(
            row=row,
            row_number=row_number,
            action=action,
            version_uuid=version_uuid,
            code=ERR_XLSX_INVALID_ACTION,
            message="Action XLSX invalide. Valeurs permises : update, ignore, archive, new.",
        )

    if not version_uuid:
        if action == "new":
            return {
                "row_number": row_number,
                "action": "new",
                "version_uuid": None,
                "blocked": False,
                "exists": False,
                "diff_count": 0,
                "diffs": [],
                "updates": {},
                "row": row,
                "warnings": [],
                "errors": [],
            }

        return _blocked_change(
            row=row,
            row_number=row_number,
            action=action,
            version_uuid=None,
            code=ERR_VERSION_UUID_MISSING,
            message="version_uuid est obligatoire pour comparer une ligne XLSX avec SQLite.",
        )

    existing_row = get_library_row_by_version_uuid(connection, version_uuid)

    if existing_row is None:
        if action in {"new", "ignore"}:
            return {
                "row_number": row_number,
                "action": "new",
                "version_uuid": version_uuid,
                "blocked": False,
                "exists": False,
                "diff_count": 0,
                "diffs": [],
                "updates": {},
                "row": row,
                "warnings": [],
                "errors": [],
            }

        return _blocked_change(
            row=row,
            row_number=row_number,
            action=action,
            version_uuid=version_uuid,
            code=ERR_VERSION_UUID_MISSING,
            message="Aucune ligne SQLite trouvée pour ce version_uuid.",
        )

    if action == "archive":
        return {
            "row_number": row_number,
            "action": "archive",
            "version_uuid": version_uuid,
            "blocked": False,
            "exists": True,
            "before": existing_row,
            "diff_count": 1 if existing_row.get("status") != "archived" else 0,
            "diffs": [
                {
                    "field": "status",
                    "sqlite_value": existing_row.get("status"),
                    "xlsx_value": "archived",
                    "protected": False,
                }
            ]
            if existing_row.get("status") != "archived"
            else [],
            "updates": {
                "status": "archived",
            }
            if existing_row.get("status") != "archived"
            else {},
            "row": row,
            "warnings": [],
            "errors": [],
        }

    protected_attempts = _detect_protected_update_attempts(
        existing_row=existing_row,
        xlsx_row=row,
    )

    updates = xlsx_row_to_update_dict(row)
    diffs = _build_diffs(
        existing_row=existing_row,
        updates=updates,
    )

    warnings: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    for field_name in protected_attempts:
        message = KoaMessage(
            code=ERR_PROTECTED_FIELD_UPDATE,
            severity="blocking",
            field=field_name,
            row_number=row_number,
            message=f"Tentative de modification d'un champ protégé : {field_name}.",
        )
        errors.append(_message_to_dict(message))

    blocked = bool(protected_attempts)

    effective_action = "update" if diffs else "ignore"
    if action == "update":
        effective_action = "update" if diffs or blocked else "ignore"
    elif action == "ignore":
        effective_action = "ignore"

    return {
        "row_number": row_number,
        "action": effective_action,
        "xlsx_action": action,
        "version_uuid": version_uuid,
        "blocked": blocked,
        "exists": True,
        "before": existing_row,
        "diff_count": len(diffs),
        "diffs": diffs,
        "updates": updates if diffs else {},
        "row": row,
        "warnings": warnings,
        "errors": errors,
    }


def _build_diffs(
    *,
    existing_row: dict[str, Any],
    updates: dict[str, Any],
) -> list[dict[str, Any]]:
    diffs: list[dict[str, Any]] = []

    protected = set(XLSX_PROTECTED_COLUMNS) | set(TECHNICAL_PROTECTED_FIELDS)

    for field_name, xlsx_value in updates.items():
        if field_name == "updated_at":
            continue

        sqlite_value = existing_row.get(field_name)

        if _normalize_compare_value(sqlite_value) == _normalize_compare_value(xlsx_value):
            continue

        diffs.append(
            {
                "field": field_name,
                "sqlite_value": sqlite_value,
                "xlsx_value": xlsx_value,
                "protected": field_name in protected,
            }
        )

    return diffs


def _detect_protected_update_attempts(
    *,
    existing_row: dict[str, Any],
    xlsx_row: dict[str, Any],
) -> list[str]:
    protected_fields = set(XLSX_PROTECTED_COLUMNS) | set(TECHNICAL_PROTECTED_FIELDS)
    attempts: list[str] = []

    for field_name in sorted(protected_fields):
        if field_name not in xlsx_row:
            continue

        incoming = _normalize_compare_value(xlsx_row.get(field_name))
        existing = _normalize_compare_value(existing_row.get(field_name))

        if incoming == "" or incoming == existing:
            continue

        attempts.append(field_name)

    return attempts


def _read_library_sheet_rows(sheet) -> list[tuple[int, dict[str, Any]]]:
    headers = _read_headers(sheet)
    rows: list[tuple[int, dict[str, Any]]] = []

    for row_number in range(2, sheet.max_row + 1):
        values = {
            header: sheet.cell(row=row_number, column=col_index).value
            for col_index, header in enumerate(headers, start=1)
            if header
        }

        if _is_empty_row(values):
            continue

        rows.append((row_number, values))

    return rows


def _read_headers(sheet) -> list[str]:
    headers: list[str] = []

    for col_index in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=col_index).value
        header = str(value).strip() if value is not None else ""
        headers.append(header)

    return headers


def _blocked_change(
    *,
    row: dict[str, Any],
    row_number: int,
    action: str,
    version_uuid: str | None,
    code: str,
    message: str,
) -> dict[str, Any]:
    error = KoaMessage(
        code=code,
        severity="blocking",
        field=None,
        row_number=row_number,
        message=message,
    )

    return {
        "row_number": row_number,
        "action": action or "blocked",
        "version_uuid": version_uuid,
        "blocked": True,
        "exists": False,
        "diff_count": 0,
        "diffs": [],
        "updates": {},
        "row": row,
        "warnings": [],
        "errors": [_message_to_dict(error)],
    }


def _message_to_dict(message: KoaMessage) -> dict[str, Any]:
    return {
        "code": message.code,
        "severity": message.severity,
        "message": message.message,
        "field": message.field,
        "row_number": message.row_number,
        "details": message.details,
    }


def _normalize_action(value: Any) -> str:
    if value is None:
        return "ignore"

    return str(value).strip().lower()


def _count_changes(changes: list[dict[str, Any]], action: str) -> int:
    return sum(
        1
        for change in changes
        if change.get("action") == action and not change.get("blocked")
    )


def _is_empty_row(row: dict[str, Any]) -> bool:
    return all(value is None or str(value).strip() == "" for value in row.values())


def _clean_string(value: Any) -> str | None:
    if value is None:
        return None

    text = str(value).strip()
    return text or None


def _normalize_compare_value(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, bool):
        return "1" if value else "0"

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)

    text = str(value).strip()

    try:
        parsed = json.loads(text)
    except (json.JSONDecodeError, TypeError):
        return text

    if isinstance(parsed, (dict, list)):
        return json.dumps(parsed, ensure_ascii=False, sort_keys=True)

    return str(parsed).strip()


def _dict_row_factory(cursor, row):
    return {
        cursor.description[index][0]: value
        for index, value in enumerate(row)
    }