# 06_GUI/koa_mediatheque_gui/koa_mediatheque/services/xlsx_import_service.py

from __future__ import annotations

import json
import sqlite3
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from koa_mediatheque.errors import (
    ERR_FILE_NOT_FOUND,
    ERR_PROTECTED_FIELD_UPDATE,
    ERR_VERSION_UUID_DUPLICATE,
    ERR_VERSION_UUID_MISSING,
    ERR_XLSX_INVALID_ACTION,
    ERR_XLSX_MISSING_SHEET,
)
from koa_mediatheque.models import ImportPreview, KoaMessage, OperationResult

try:
    from koa_mediatheque.constants import XLSX_PROTECTED_COLUMNS, TECHNICAL_PROTECTED_FIELDS
except ImportError:  # pragma: no cover - defensive fallback for isolated tests
    XLSX_PROTECTED_COLUMNS = [
        "media_uuid",
        "version_uuid",
        "filename",
        "original_path",
        "storage_path",
        "sha256",
        "filesize",
        "mimetype",
        "updated_at",
    ]
    TECHNICAL_PROTECTED_FIELDS = [
        "sha256",
        "filesize",
        "mimetype",
        "filename",
        "extension",
        "original_path",
        "storage_path",
    ]


LIBRARY_SHEET_NAME = "Library"

ALLOWED_ACTIONS = {"update", "ignore", "archive", "new"}

JSON_DISPLAY_FIELDS = {
    "collections": "collections_json",
    "tags": "tags_json",
    "relations": "relations_json",
    "content_flags": "content_flags_json",
}

PROTECTED_FIELDS_NORMAL = set(XLSX_PROTECTED_COLUMNS) | set(TECHNICAL_PROTECTED_FIELDS)
PROTECTED_FIELDS_NORMAL.add("updated_at")

INTERNAL_XLSX_COLUMNS = {"xlsx_action", "action"}


def preview_xlsx_import(
    connection: sqlite3.Connection,
    xlsx_path: str | Path,
    *,
    mode: str = "normal",
) -> ImportPreview:
    source_path = Path(xlsx_path)
    import_uuid = str(uuid4())

    changes: list[dict[str, Any]] = []
    warnings: list[KoaMessage] = []
    errors: list[KoaMessage] = []

    if not source_path.exists():
        return ImportPreview(
            import_uuid=import_uuid,
            source_path=str(source_path),
            rows_total=0,
            rows_update=0,
            rows_new=0,
            rows_archive=0,
            rows_ignore=0,
            rows_blocked=1,
            changes=[],
            warnings=[],
            errors=[
                _message(
                    ERR_FILE_NOT_FOUND,
                    "blocking",
                    f"XLSX file not found: {source_path}",
                    field="xlsx_path",
                )
            ],
        )

    try:
        workbook = load_workbook(source_path, data_only=True)
    except Exception as exc:
        return ImportPreview(
            import_uuid=import_uuid,
            source_path=str(source_path),
            rows_total=0,
            rows_update=0,
            rows_new=0,
            rows_archive=0,
            rows_ignore=0,
            rows_blocked=1,
            changes=[],
            warnings=[],
            errors=[
                _message(
                    "ERR_XLSX_READ_FAILED",
                    "blocking",
                    str(exc),
                    field="xlsx_path",
                )
            ],
        )

    if LIBRARY_SHEET_NAME not in workbook.sheetnames:
        return ImportPreview(
            import_uuid=import_uuid,
            source_path=str(source_path),
            rows_total=0,
            rows_update=0,
            rows_new=0,
            rows_archive=0,
            rows_ignore=0,
            rows_blocked=1,
            changes=[],
            warnings=[],
            errors=[
                _message(
                    ERR_XLSX_MISSING_SHEET,
                    "blocking",
                    "Workbook must contain a Library sheet.",
                    field=LIBRARY_SHEET_NAME,
                )
            ],
        )

    rows = _read_library_rows(workbook[LIBRARY_SHEET_NAME])
    table_columns = _table_columns(connection, "library_rows")

    for index, row in enumerate(rows, start=2):
        action = _normalize_action(row.get("xlsx_action", row.get("action")))
        version_uuid = _clean_string(row.get("version_uuid"))

        if action not in ALLOWED_ACTIONS:
            row_error = _message(
                ERR_XLSX_INVALID_ACTION,
                "blocking",
                f"Invalid XLSX action: {action or '<blank>'}",
                field="xlsx_action",
                row_number=index,
                details={"allowed": sorted(ALLOWED_ACTIONS)},
            )
            errors.append(row_error)
            changes.append(
                _change(
                    row_number=index,
                    action=action,
                    version_uuid=version_uuid,
                    blocked=True,
                    errors=[row_error],
                    raw_row=row,
                )
            )
            continue

        if action == "ignore":
            changes.append(
                _change(
                    row_number=index,
                    action=action,
                    version_uuid=version_uuid,
                    blocked=False,
                    raw_row=row,
                )
            )
            continue

        if action in {"update", "archive"} and not version_uuid:
            row_error = _message(
                ERR_VERSION_UUID_MISSING,
                "blocking",
                "version_uuid is required for update and archive actions.",
                field="version_uuid",
                row_number=index,
            )
            errors.append(row_error)
            changes.append(
                _change(
                    row_number=index,
                    action=action,
                    version_uuid=version_uuid,
                    blocked=True,
                    errors=[row_error],
                    raw_row=row,
                )
            )
            continue

        if action in {"update", "archive"}:
            existing = _get_row_by_version_uuid(connection, version_uuid)

            if existing is None:
                row_error = _message(
                    ERR_VERSION_UUID_MISSING,
                    "blocking",
                    "version_uuid was not found in SQLite.",
                    field="version_uuid",
                    row_number=index,
                    details={"version_uuid": version_uuid},
                )
                errors.append(row_error)
                changes.append(
                    _change(
                        row_number=index,
                        action=action,
                        version_uuid=version_uuid,
                        blocked=True,
                        errors=[row_error],
                        raw_row=row,
                    )
                )
                continue

            protected_warnings = _protected_field_warnings(
                row=row,
                existing=existing,
                row_number=index,
                mode=mode,
            )
            warnings.extend(protected_warnings)

            updates = _xlsx_row_to_updates(
                row,
                table_columns=table_columns,
                mode=mode,
                include_protected=False,
            )

            changes.append(
                _change(
                    row_number=index,
                    action=action,
                    version_uuid=version_uuid,
                    media_uuid=_clean_string(existing.get("media_uuid")),
                    blocked=False,
                    warnings=protected_warnings,
                    updates=updates,
                    before=existing,
                    raw_row=row,
                )
            )
            continue

        if action == "new":
            new_version_uuid = version_uuid or str(uuid4())
            existing = _get_row_by_version_uuid(connection, new_version_uuid)

            if existing is not None:
                row_error = _message(
                    ERR_VERSION_UUID_DUPLICATE,
                    "blocking",
                    "version_uuid already exists in SQLite.",
                    field="version_uuid",
                    row_number=index,
                    details={"version_uuid": new_version_uuid},
                )
                errors.append(row_error)
                changes.append(
                    _change(
                        row_number=index,
                        action=action,
                        version_uuid=new_version_uuid,
                        blocked=True,
                        errors=[row_error],
                        raw_row=row,
                    )
                )
                continue

            row_data = _xlsx_row_to_insert_data(
                row,
                table_columns=table_columns,
                version_uuid=new_version_uuid,
            )
            row_errors = _validate_new_row(row_data, row_number=index)

            if row_errors:
                errors.extend(row_errors)
                changes.append(
                    _change(
                        row_number=index,
                        action=action,
                        version_uuid=new_version_uuid,
                        media_uuid=_clean_string(row_data.get("media_uuid")),
                        blocked=True,
                        errors=row_errors,
                        raw_row=row,
                    )
                )
                continue

            changes.append(
                _change(
                    row_number=index,
                    action=action,
                    version_uuid=new_version_uuid,
                    media_uuid=_clean_string(row_data.get("media_uuid")),
                    blocked=False,
                    updates=row_data,
                    raw_row=row,
                )
            )

    return ImportPreview(
        import_uuid=import_uuid,
        source_path=str(source_path),
        rows_total=len(rows),
        rows_update=_count_unblocked(changes, "update"),
        rows_new=_count_unblocked(changes, "new"),
        rows_archive=_count_unblocked(changes, "archive"),
        rows_ignore=_count_unblocked(changes, "ignore"),
        rows_blocked=sum(1 for change in changes if change.get("blocked")),
        changes=changes,
        warnings=warnings,
        errors=errors,
    )


def apply_xlsx_import(
    connection: sqlite3.Connection,
    xlsx_path: str | Path,
    *,
    backup_dir: str | Path,
    actor: str = "local_user",
    mode: str = "normal",
) -> OperationResult:
    preview = preview_xlsx_import(connection, xlsx_path, mode=mode)

    if preview.errors:
        _write_import_log(
            connection,
            preview=preview,
            mode=mode,
            status="blocked",
        )

        return OperationResult(
            success=False,
            operation="apply_xlsx_import",
            result="blocked",
            entity_type="xlsx_import",
            entity_uuid=preview.import_uuid,
            path=str(xlsx_path),
            data=_preview_to_data(preview),
            warnings=preview.warnings,
            errors=preview.errors,
        )

    if mode == "dry_run":
        return OperationResult(
            success=True,
            operation="apply_xlsx_import",
            result="dry_run",
            entity_type="xlsx_import",
            entity_uuid=preview.import_uuid,
            path=str(xlsx_path),
            data=_preview_to_data(preview),
            warnings=preview.warnings,
            errors=[],
        )

    backup_path = _backup_connection(connection, backup_dir)
    applied: list[dict[str, Any]] = []

    try:
        for change in preview.changes:
            if change.get("blocked"):
                continue

            action = change["action"]

            if action == "ignore":
                continue

            if action == "update":
                _apply_update(connection, change)
                applied.append(change)
                _write_audit_log(
                    connection,
                    change,
                    actor=actor,
                    action="xlsx_import_applied",
                )
                continue

            if action == "archive":
                _apply_archive(connection, change)
                applied.append(change)
                _write_audit_log(
                    connection,
                    change,
                    actor=actor,
                    action="library_row_archived",
                )
                continue

            if action == "new":
                _apply_new(connection, change)
                applied.append(change)
                _write_audit_log(
                    connection,
                    change,
                    actor=actor,
                    action="library_row_inserted",
                )
                continue

        _write_import_log(
            connection,
            preview=preview,
            mode=mode,
            status="imported",
        )

        connection.commit()

        return OperationResult(
            success=True,
            operation="apply_xlsx_import",
            result="imported",
            entity_type="xlsx_import",
            entity_uuid=preview.import_uuid,
            path=str(xlsx_path),
            data={
                **_preview_to_data(preview),
                "backup_path": str(backup_path),
                "applied_rows": len(applied),
            },
            warnings=preview.warnings,
            errors=[],
        )
    except Exception as exc:
        connection.rollback()

        return OperationResult(
            success=False,
            operation="apply_xlsx_import",
            result="failed",
            entity_type="xlsx_import",
            entity_uuid=preview.import_uuid,
            path=str(xlsx_path),
            data={
                **_preview_to_data(preview),
                "backup_path": str(backup_path),
            },
            warnings=preview.warnings,
            errors=[
                _message(
                    "ERR_XLSX_IMPORT_FAILED",
                    "blocking",
                    str(exc),
                    field="xlsx_path",
                )
            ],
        )


def _read_library_rows(worksheet) -> list[dict[str, Any]]:
    header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)

    if not header_values:
        return []

    headers = [str(value).strip() if value is not None else "" for value in header_values]
    rows: list[dict[str, Any]] = []

    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = {
            headers[index]: _clean_cell(value)
            for index, value in enumerate(values)
            if index < len(headers) and headers[index]
        }

        if _is_empty_row(row):
            continue

        rows.append(row)

    return rows


def _xlsx_row_to_updates(
    row: dict[str, Any],
    *,
    table_columns: set[str],
    mode: str,
    include_protected: bool,
) -> dict[str, Any]:
    updates: dict[str, Any] = {}

    for key, value in row.items():
        if key in INTERNAL_XLSX_COLUMNS:
            continue

        if key in JSON_DISPLAY_FIELDS:
            target = JSON_DISPLAY_FIELDS[key]
            if target in table_columns:
                updates[target] = _to_json_array_text(value)
            continue

        if key not in table_columns:
            continue

        if not include_protected and _is_protected_field(key, mode):
            continue

        if value is None:
            continue

        updates[key] = value

    return updates


def _xlsx_row_to_insert_data(
    row: dict[str, Any],
    *,
    table_columns: set[str],
    version_uuid: str,
) -> dict[str, Any]:
    data = _xlsx_row_to_updates(
        row,
        table_columns=table_columns,
        mode="repair",
        include_protected=True,
    )

    data["version_uuid"] = version_uuid

    if "media_uuid" in table_columns and not _clean_string(data.get("media_uuid")):
        data["media_uuid"] = str(uuid4())

    for source, target in JSON_DISPLAY_FIELDS.items():
        if target in table_columns and target not in data:
            data[target] = _to_json_array_text(row.get(source))

    return data


def _validate_new_row(row_data: dict[str, Any], *, row_number: int) -> list[KoaMessage]:
    errors: list[KoaMessage] = []

    for field_name in ("media_uuid", "version_uuid", "title", "original_path", "filename"):
        if not _clean_string(row_data.get(field_name)):
            errors.append(
                _message(
                    "ERR_REQUIRED_FIELD",
                    "blocking",
                    f"{field_name} is required for new rows.",
                    field=field_name,
                    row_number=row_number,
                )
            )

    return errors


def _apply_update(connection: sqlite3.Connection, change: dict[str, Any]) -> None:
    updates = dict(change.get("updates") or {})

    updates.pop("updated_at", None)
    updates.pop("created_at", None)

    if not updates:
        return

    _update_library_row(connection, change["version_uuid"], updates)


def _apply_archive(connection: sqlite3.Connection, change: dict[str, Any]) -> None:
    _update_library_row(
        connection,
        change["version_uuid"],
        {"status": "archived"},
    )


def _apply_new(connection: sqlite3.Connection, change: dict[str, Any]) -> None:
    row_data = dict(change.get("updates") or {})

    if not row_data:
        return

    table_columns = _table_columns(connection, "library_rows")
    row_data = {
        key: value
        for key, value in row_data.items()
        if key in table_columns and key != "id"
    }

    if not row_data:
        return

    columns = list(row_data.keys())
    placeholders = ", ".join("?" for _ in columns)

    connection.execute(
        f"""
        INSERT INTO library_rows ({", ".join(columns)})
        VALUES ({placeholders})
        """,
        [row_data[column] for column in columns],
    )


def _update_library_row(
    connection: sqlite3.Connection,
    version_uuid: str,
    updates: dict[str, Any],
) -> None:
    if not updates:
        return

    table_columns = _table_columns(connection, "library_rows")

    safe_updates = {
        key: value
        for key, value in updates.items()
        if key in table_columns and key not in {"id", "version_uuid"}
    }

    if not safe_updates:
        return

    assignments = ", ".join(f"{column} = ?" for column in safe_updates)
    values = list(safe_updates.values()) + [version_uuid]

    connection.execute(
        f"""
        UPDATE library_rows
        SET {assignments}
        WHERE version_uuid = ?
        """,
        values,
    )


def _backup_connection(connection: sqlite3.Connection, backup_dir: str | Path) -> Path:
    backup_root = Path(backup_dir)
    backup_root.mkdir(parents=True, exist_ok=True)

    backup_path = backup_root / f"koa_mediatheque_xlsx_import_{uuid4()}.sqlite"

    destination = sqlite3.connect(backup_path)
    try:
        connection.backup(destination)
        destination.commit()
    finally:
        destination.close()

    return backup_path


def _write_import_log(
    connection: sqlite3.Connection,
    *,
    preview: ImportPreview,
    mode: str,
    status: str,
) -> None:
    if not _table_exists(connection, "xlsx_import_log"):
        return

    _insert_partial(
        connection,
        "xlsx_import_log",
        {
            "import_uuid": preview.import_uuid,
            "xlsx_path": preview.source_path,
            "mode": mode,
            "rows_total": preview.rows_total,
            "rows_update": preview.rows_update,
            "rows_new": preview.rows_new,
            "rows_archive": preview.rows_archive,
            "rows_ignore": preview.rows_ignore,
            "rows_blocked": preview.rows_blocked,
            "status": status,
            "report_json": _json_dumps(_preview_to_data(preview)),
        },
    )


def _write_audit_log(
    connection: sqlite3.Connection,
    change: dict[str, Any],
    *,
    actor: str,
    action: str,
) -> None:
    if not _table_exists(connection, "audit_log"):
        return

    _insert_partial(
        connection,
        "audit_log",
        {
            "action": action,
            "entity_type": "library_row",
            "entity_uuid": change.get("version_uuid"),
            "before_json": _json_dumps(change.get("before") or {}),
            "after_json": _json_dumps(
                {
                    "action": change.get("action"),
                    "updates": change.get("updates") or {},
                    "raw_row": change.get("raw_row") or {},
                }
            ),
            "actor": actor,
            "note": f"XLSX import action={change.get('action')}",
        },
    )


def _insert_partial(
    connection: sqlite3.Connection,
    table_name: str,
    payload: dict[str, Any],
) -> None:
    columns = _table_columns(connection, table_name)
    filtered = {key: value for key, value in payload.items() if key in columns}

    if not filtered:
        return

    column_sql = ", ".join(filtered)
    placeholders = ", ".join("?" for _ in filtered)

    connection.execute(
        f"""
        INSERT INTO {table_name} ({column_sql})
        VALUES ({placeholders})
        """,
        list(filtered.values()),
    )


def _get_row_by_version_uuid(
    connection: sqlite3.Connection,
    version_uuid: str | None,
) -> dict[str, Any] | None:
    if not version_uuid:
        return None

    cursor = connection.execute(
        "SELECT * FROM library_rows WHERE version_uuid = ? LIMIT 1",
        (version_uuid,),
    )
    row = cursor.fetchone()

    if row is None:
        return None

    if isinstance(row, sqlite3.Row):
        return dict(row)

    return {cursor.description[index][0]: value for index, value in enumerate(row)}


def _table_exists(connection: sqlite3.Connection, table_name: str) -> bool:
    row = connection.execute(
        """
        SELECT 1
        FROM sqlite_master
        WHERE type = 'table'
          AND name = ?
        LIMIT 1
        """,
        (table_name,),
    ).fetchone()

    return row is not None


def _table_columns(connection: sqlite3.Connection, table_name: str) -> set[str]:
    rows = connection.execute(f"PRAGMA table_info({table_name})").fetchall()
    return {row[1] for row in rows}


def _protected_field_warnings(
    *,
    row: dict[str, Any],
    existing: dict[str, Any],
    row_number: int,
    mode: str,
) -> list[KoaMessage]:
    if mode != "normal":
        return []

    warnings: list[KoaMessage] = []

    for field_name in sorted(PROTECTED_FIELDS_NORMAL):
        if field_name not in row:
            continue

        incoming = row.get(field_name)

        if incoming is None:
            continue

        current = existing.get(field_name)

        if _normalize_compare_value(incoming) == _normalize_compare_value(current):
            continue

        warnings.append(
            _message(
                ERR_PROTECTED_FIELD_UPDATE,
                "warning",
                f"Protected field ignored in normal XLSX import: {field_name}",
                field=field_name,
                row_number=row_number,
                details={
                    "before": current,
                    "incoming": incoming,
                },
            )
        )

    return warnings


def _is_protected_field(field_name: str, mode: str) -> bool:
    if mode == "repair":
        return False

    return field_name in PROTECTED_FIELDS_NORMAL


def _to_json_array_text(value: Any) -> str:
    if value is None:
        return "[]"

    if isinstance(value, list):
        return _json_dumps([_clean_item(item) for item in value if _clean_item(item)])

    text = str(value).strip()

    if not text:
        return "[]"

    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return _json_dumps([_clean_item(item) for item in parsed if _clean_item(item)])
        return _json_dumps([parsed])
    except json.JSONDecodeError:
        pass

    if ";" in text:
        values = [_clean_item(part) for part in text.split(";")]
    else:
        values = [_clean_item(text)]

    return _json_dumps([value for value in values if value])


def _change(
    *,
    row_number: int,
    action: str,
    version_uuid: str | None,
    media_uuid: str | None = None,
    blocked: bool,
    warnings: list[KoaMessage] | None = None,
    errors: list[KoaMessage] | None = None,
    updates: dict[str, Any] | None = None,
    before: dict[str, Any] | None = None,
    raw_row: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "row_number": row_number,
        "action": action,
        "version_uuid": version_uuid,
        "media_uuid": media_uuid,
        "blocked": blocked,
        "warnings": warnings or [],
        "errors": errors or [],
        "updates": updates or {},
        "before": before or {},
        "raw_row": raw_row or {},
    }


def _preview_to_data(preview: ImportPreview) -> dict[str, Any]:
    return {
        "import_uuid": preview.import_uuid,
        "source_path": preview.source_path,
        "rows_total": preview.rows_total,
        "rows_update": preview.rows_update,
        "rows_new": preview.rows_new,
        "rows_archive": preview.rows_archive,
        "rows_ignore": preview.rows_ignore,
        "rows_blocked": preview.rows_blocked,
        "changes": [_change_to_serializable(change) for change in preview.changes],
        "warnings": [_message_to_dict(message) for message in preview.warnings],
        "errors": [_message_to_dict(message) for message in preview.errors],
    }


def _change_to_serializable(change: dict[str, Any]) -> dict[str, Any]:
    return {
        key: [_message_to_dict(item) for item in value]
        if key in {"warnings", "errors"}
        else value
        for key, value in change.items()
    }


def _message(
    code: str,
    severity: str,
    message: str,
    *,
    field: str | None = None,
    row_number: int | None = None,
    details: dict[str, Any] | None = None,
) -> KoaMessage:
    return KoaMessage(
        code=code,
        severity=severity,
        message=message,
        field=field,
        row_number=row_number,
        details=details or {},
    )


def _message_to_dict(message: KoaMessage) -> dict[str, Any]:
    return {
        "code": message.code,
        "severity": message.severity,
        "message": message.message,
        "field": message.field,
        "row_number": message.row_number,
        "details": message.details,
    }


def _normalize_action(value: Any) -> str:
    if value is None:
        return "ignore"

    return str(value).strip().lower()


def _count_unblocked(changes: list[dict[str, Any]], action: str) -> int:
    return sum(1 for change in changes if change.get("action") == action and not change.get("blocked"))


def _is_empty_row(row: dict[str, Any]) -> bool:
    return all(value is None or str(value).strip() == "" for value in row.values())


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, str):
        text = value.strip()
        return text if text else None

    return value


def _clean_string(value: Any) -> str | None:
    if value is None:
        return None

    text = str(value).strip()
    return text or None


def _clean_item(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def _normalize_compare_value(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def _json_dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)